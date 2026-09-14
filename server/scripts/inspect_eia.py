"""Inventory every supplied workbook, including header locations and sample data."""
import json
from pathlib import Path
import openpyxl

root = Path(__file__).resolve().parents[2]
inventory = []
for folder in ('data', 'f923_2025er'):
    for path in sorted((root / folder).glob('*.xlsx')):
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        for ws in wb:
            rows = list(ws.iter_rows(min_row=1, max_row=min(ws.max_row, 9), values_only=True))
            inventory.append({'file': str(path.relative_to(root)), 'sheet': ws.title,
                              'rows': ws.max_row, 'columns': ws.max_column, 'head': rows})
        wb.close()
out = root / 'server' / 'import-output'
out.mkdir(exist_ok=True)
(out / 'inventory.json').write_text(json.dumps(inventory, default=str, indent=2))
for s in inventory:
    print(s['file'], '|', s['sheet'], '|', s['rows'], 'x', s['columns'])
